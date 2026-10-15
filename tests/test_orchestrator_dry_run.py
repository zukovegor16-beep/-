from pathlib import Path

from openpyxl import Workbook, load_workbook

from tiktok_publisher.config import AppSettings
from tiktok_publisher.orchestrator import Orchestrator


def test_dry_run_sequential(tmp_path: Path):
    xlsx = tmp_path / "acc.xlsx"
    folder = tmp_path / "acc01"
    folder.mkdir()
    (folder / "v1.mp4").write_bytes(b"1")
    (folder / "v2.mp4").write_bytes(b"2")

    wb = Workbook()
    ws = wb.active
    ws.title = "Аккаунты"
    ws.append(
        [
            "account_no",
            "account_name",
            "video_folder",
            "dolphin_profile_id",
            "videos_per_run",
            "caption_variants",
            "active",
            "status",
            "last_run_error",
        ]
    )
    ws.append([1, "A", str(folder), "111", 2, "text1|||text2", "да", "", ""])
    wb.save(xlsx)

    data_dir = tmp_path / "data"
    settings = AppSettings(
        dry_run=True,
        paths={"data_dir": str(data_dir), "content_root": str(tmp_path)},
        excel={
            "registry_path": str(xlsx),
            "sheet_accounts": "Аккаунты",
        },
        orchestrator={
            "speed_preset": "custom",
            "min_video_bytes": 1,
            "delay_between_posts_sec": {"min": 0, "max": 0},
            "delay_after_error_sec": 0,
        },
    )

    orch = Orchestrator(settings)
    orch.start()
    orch.wait(timeout=30)
    assert not orch.is_running
    recs = orch.store.publications_for_run(orch.current_run_id or "")
    assert len(recs) == 2
    assert all(r.status == "dry_run" for r in recs)
    assert all(r.status_label == "тест без публикации" for r in recs)
    assert not (folder / "v1.mp4").exists()
    assert (folder / "posted" / "v1.mp4").exists()
    wb = load_workbook(xlsx)
    row = list(wb["Аккаунты"].iter_rows(min_row=2, max_row=2, values_only=True))[0]
    assert row[7] == "тест без публикации"
