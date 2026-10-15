from pathlib import Path

import yaml
from openpyxl import Workbook

from tiktok_publisher.config import AppSettings
from tiktok_publisher.excel_registry import ExcelRegistry
from tiktok_publisher.geo_catalog import load_geo_catalog, resolve_geo
from tiktok_publisher.onboarding_orchestrator import OnboardingOrchestrator
from tiktok_publisher.outcomes import OutcomeCode
from tiktok_publisher.warmup import run_warmup
import asyncio


def test_geo_catalog_load():
    catalog = load_geo_catalog(Path("config/geo_catalog.yaml"))
    assert "US" in catalog
    spec = resolve_geo(AppSettings(), "us")
    assert spec.timezone


def test_onboard_dry_run_promote(tmp_path: Path):
    xlsx = tmp_path / "acc.xlsx"
    wb = Workbook()
    ws_acc = wb.active
    ws_acc.title = "Аккаунты"
    ws_acc.append(
        [
            "account_no",
            "account_name",
            "video_folder",
            "dolphin_profile_id",
            "videos_per_run",
            "caption_variants",
            "active",
            "geo",
        ]
    )
    ob = wb.create_sheet("Онбординг")
    ob.append(
        [
            "geo",
            "tiktok_login",
            "tiktok_password_current",
            "tiktok_password_new",
            "email",
            "email_password",
            "avatar_path",
            "proxy",
            "dolphin_profile_id",
            "onboard_status",
            "onboard_error",
        ]
    )
    ob.append(["US", "user1", "pass1", "pass2", "", "", "", "", "", "", ""])
    wb.save(xlsx)

    geo_yaml = tmp_path / "geo.yaml"
    geo_yaml.write_text(
        yaml.safe_dump(
            {
                "US": {
                    "latitude": 40.7,
                    "longitude": -74.0,
                    "timezone": "America/New_York",
                    "locale": "en-US",
                }
            }
        ),
        encoding="utf-8",
    )

    settings = AppSettings(
        dry_run=True,
        excel={
            "registry_path": str(xlsx),
            "sheet_accounts": "Аккаунты",
            "sheet_onboarding": "Онбординг",
        },
        geo_catalog_path=str(geo_yaml),
    )
    ob_orch = OnboardingOrchestrator(settings)
    ob_orch.start()
    ob_orch.wait(30)

    reg = ExcelRegistry(settings)
    assert len(reg.read_accounts()) == 0
    wb2 = __import__("openpyxl").load_workbook(xlsx, data_only=True)
    ws = wb2["Аккаунты"]
    assert ws.max_row == 2
    assert ws.cell(2, 1).value == "user1"
    assert str(ws.cell(2, 7).value).lower() in {"нет", "no", "false"}
    onboard_rows = reg.read_onboarding_rows()
    assert onboard_rows[0].onboard_status == "готов"


def test_warmup_dry_run():
    settings = AppSettings(dry_run=True)
    loop = asyncio.new_event_loop()
    try:
        wr = run_warmup(settings, None, loop)
        assert wr.outcome.code == OutcomeCode.dry_run
    finally:
        loop.close()


def test_session_order_helpers():
    from tiktok_publisher.orchestrator import Orchestrator

    s = AppSettings(session={"order": "warmup_then_publish"})
    o = Orchestrator(s)
    assert o._needs_warmup_before()
    assert not o._needs_warmup_after()
