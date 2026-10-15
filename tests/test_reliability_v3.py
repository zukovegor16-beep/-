from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from tiktok_publisher.asset_fingerprint import file_fingerprint
from tiktok_publisher.config import AppSettings
from tiktok_publisher.excel_registry import ExcelRegistry, _atomic_save_workbook
from tiktok_publisher.orchestrator import Orchestrator
from tiktok_publisher.posting_window import in_posting_window
from tiktok_publisher.preflight import run_preflight
from tiktok_publisher.speed_presets import apply_speed_preset
from tiktok_publisher.state import StateStore


def test_duplicate_skipped_second_run(tmp_path: Path):
    folder = tmp_path / "acc01"
    folder.mkdir()
    video = folder / "v1.mp4"
    video.write_bytes(b"content123")

    xlsx = tmp_path / "acc.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Аккаунты"
    ws.append(
        [
            "account_no",
            "account_name",
            "video_folder",
            "dolphin_profile_id",
            "videos_per_run",
            "caption_variants",
            "active",
        ]
    )
    ws.append([1, "A", str(folder), "111", 1, "t1", "да"])
    wb.save(xlsx)

    settings = AppSettings(
        dry_run=True,
        paths={"data_dir": str(tmp_path / "data"), "content_root": str(tmp_path)},
        excel={"registry_path": str(xlsx), "sheet_accounts": "Аккаунты"},
        orchestrator={
            "speed_preset": "custom",
            "min_video_bytes": 1,
            "skip_duplicate_counts_toward_limit": True,
            "delay_between_posts_sec": {"min": 0, "max": 0},
            "delay_after_error_sec": 0,
        },
    )
    store = StateStore(settings.db_path)
    fp = file_fingerprint(video)
    store.record_published_asset("1", fp, video.name, str(video), "ok", "prev", "2020-01-01T00:00:00")

    orch = Orchestrator(settings)
    orch.start()
    orch.wait(timeout=30)
    recs = store.publications_for_run(orch.current_run_id or "")
    assert len(recs) == 1
    assert recs[0].status == "skipped_duplicate"
    assert (folder / "posted" / "v1.mp4").exists()


def test_preflight_duplicate_profile(tmp_path: Path):
    xlsx = tmp_path / "acc.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Аккаунты"
    ws.append(
        [
            "account_no",
            "account_name",
            "video_folder",
            "dolphin_profile_id",
            "videos_per_run",
            "caption_variants",
            "active",
        ]
    )
    f = tmp_path / "a"
    f.mkdir()
    ws.append([1, "A", str(f), "same", 1, "c", "да"])
    ws.append([2, "B", str(f), "same", 1, "c", "да"])
    wb.save(xlsx)

    settings = AppSettings(
        dry_run=True,
        excel={"registry_path": str(xlsx), "sheet_accounts": "Аккаунты"},
    )
    pf = run_preflight(settings)
    assert not pf.ok
    assert any(i.code == "dup_profile" for i in pf.issues)


def test_posting_window_reject():
    orch = apply_speed_preset(
        AppSettings(
            orchestrator={
                "speed_preset": "custom",
                "posting_window": {"start": "23:00", "end": "23:01"},
                "timezone": "UTC",
            }
        )
    )
    from datetime import datetime
    from zoneinfo import ZoneInfo

    inside = datetime(2026, 1, 1, 23, 0, 30, tzinfo=ZoneInfo("UTC"))
    outside = datetime(2026, 1, 1, 12, 0, 0, tzinfo=ZoneInfo("UTC"))
    assert in_posting_window(orch, inside)
    assert not in_posting_window(orch, outside)


def test_atomic_save_workbook(tmp_path: Path):
    dest = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active["A1"] = "x"
    _atomic_save_workbook(wb, dest)
    assert dest.is_file()
    wb2 = Workbook()
    wb2.active["A1"] = "y"
    with patch("tiktok_publisher.excel_registry.Path.replace", side_effect=OSError("fail")):
        try:
            _atomic_save_workbook(wb2, dest)
        except OSError:
            pass
    reg = ExcelRegistry(
        AppSettings(excel={"registry_path": str(dest), "sheet_accounts": "Sheet"})
    )
    assert dest.is_file()
