from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

from openpyxl import Workbook, load_workbook

from tiktok_publisher.config import AppSettings
from tiktok_publisher.excel_pending import sync_pending_excel
from tiktok_publisher.excel_registry import ExcelRegistry
from tiktok_publisher.orchestrator import Orchestrator
from tiktok_publisher.outcomes import OutcomeCode
from tiktok_publisher.publish_result import PublishResult


def _accounts_workbook(xlsx: Path, folder: Path, *, videos_per_run: int = 2) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Аккаунты"
    ws.append(
        [
            "account_no",
            "account_name",
            "video_folder",
            "dolphin_profile_id",
            "videos_per_run",
            "caption_variants",
            "active",
            "status",
            "last_run_error",
        ]
    )
    ws.append(
        [1, "A", str(folder), "111", videos_per_run, "t1|||t2", "да", "", ""]
    )
    wb.save(xlsx)


def test_dry_run_log_sheet_ru_status(tmp_path: Path):
    xlsx = tmp_path / "acc.xlsx"
    folder = tmp_path / "acc01"
    folder.mkdir()
    (folder / "v1.mp4").write_bytes(b"x")
    (folder / "v2.mp4").write_bytes(b"y")
    _accounts_workbook(xlsx, folder)

    settings = AppSettings(
        dry_run=True,
        paths={"data_dir": str(tmp_path / "data"), "content_root": str(tmp_path)},
        excel={"registry_path": str(xlsx), "sheet_accounts": "Аккаунты"},
        session={"order": "publish_only"},
        orchestrator={
            "speed_preset": "custom",
            "min_video_bytes": 1,
            "delay_between_posts_sec": {"min": 0, "max": 0},
            "delay_after_error_sec": 0,
        },
    )
    orch = Orchestrator(settings)
    orch.start()
    orch.wait(timeout=30)

    wb = load_workbook(xlsx, data_only=True)
    assert "Log" in wb.sheetnames
    log_rows = list(wb["Log"].iter_rows(min_row=2, values_only=True))
    assert len(log_rows) == 2
    assert all(r[5] == "тест без публикации" for r in log_rows)


def test_partial_run_account_summary_ru(tmp_path: Path):
    xlsx = tmp_path / "acc.xlsx"
    folder = tmp_path / "acc01"
    folder.mkdir()
    (folder / "v1.mp4").write_bytes(b"1")
    (folder / "v2.mp4").write_bytes(b"2")
    _accounts_workbook(xlsx, folder)

    settings = AppSettings(
        dry_run=False,
        paths={"data_dir": str(tmp_path / "data"), "content_root": str(tmp_path)},
        excel={"registry_path": str(xlsx), "sheet_accounts": "Аккаунты"},
        session={"order": "publish_only"},
        orchestrator={
            "speed_preset": "custom",
            "min_video_bytes": 1,
            "delay_between_posts_sec": {"min": 0, "max": 0},
            "delay_after_error_sec": 0,
        },
    )

    results = [
        PublishResult(True, OutcomeCode.published, ""),
        PublishResult(False, OutcomeCode.not_published, "confirm fail"),
    ]
    publish = AsyncMock(side_effect=results)
    handle = MagicMock()
    handle.is_healthy = AsyncMock(return_value=True)

    orch = Orchestrator(settings)
    with (
        patch.object(orch.dolphin, "start_profile"),
        patch.object(orch.dolphin, "stop_profile"),
        patch("tiktok_publisher.orchestrator.open_browser", return_value=handle),
        patch("tiktok_publisher.orchestrator.close_browser"),
        patch("tiktok_publisher.orchestrator.publish_video_async", publish),
    ):
        orch.start()
        orch.wait(timeout=30)

    wb = load_workbook(xlsx, data_only=True)
    acc_status = wb["Аккаунты"].cell(row=2, column=8).value
    assert "частично" in str(acc_status)
    log_rows = list(wb["Log"].iter_rows(min_row=2, values_only=True))
    labels = {r[5] for r in log_rows}
    assert "опубликовано" in labels
    assert "нет публикации" in labels


def test_pending_excel_sync_restores_ru_status(tmp_path: Path):
    xlsx = tmp_path / "acc.xlsx"
    folder = tmp_path / "acc01"
    folder.mkdir()
    (folder / "v1.mp4").write_bytes(b"x")
    _accounts_workbook(xlsx, folder, videos_per_run=1)

    settings = AppSettings(
        dry_run=True,
        paths={"data_dir": str(tmp_path / "data"), "content_root": str(tmp_path)},
        excel={"registry_path": str(xlsx), "sheet_accounts": "Аккаунты"},
        session={"order": "publish_only"},
        orchestrator={
            "speed_preset": "custom",
            "min_video_bytes": 1,
            "delay_between_posts_sec": {"min": 0, "max": 0},
            "delay_after_error_sec": 0,
        },
    )

    orch = Orchestrator(settings)
    with patch.object(ExcelRegistry, "append_log_row", return_value=False):
        orch.start()
        orch.wait(timeout=30)

    assert orch._excel_pending.pending_count() >= 1
    applied, failed = sync_pending_excel(settings)
    assert applied >= 1
    assert failed == 0

    wb = load_workbook(xlsx, data_only=True)
    log_rows = list(wb["Log"].iter_rows(min_row=2, values_only=True))
    assert log_rows[0][5] == "тест без публикации"
