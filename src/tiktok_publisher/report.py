from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from tiktok_publisher.queue import list_pending_videos
from tiktok_publisher.config import AppSettings
from tiktok_publisher.excel_registry import AccountRow
from tiktok_publisher.state import PublicationRecord, StateStore


def write_run_report(
    path: Path,
    run_id: str,
    records: list[PublicationRecord],
    accounts: list[AccountRow],
    settings: AppSettings,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    pub_rows = [
        {
            "account_no": r.account_no,
            "account_name": r.account_name,
            "file": r.file_path,
            "caption_used": r.caption_used,
            "status": r.status,
            "error": r.error_message,
            "started_at": r.started_at,
            "finished_at": r.finished_at,
        }
        for r in records
    ]
    df_pub = pd.DataFrame(pub_rows)

    summary = []
    for acc in accounts:
        ok = sum(1 for r in records if r.account_no == acc.account_no and r.status == "ok")
        err = sum(1 for r in records if r.account_no == acc.account_no and r.status == "error")
        pending = len(list_pending_videos(acc.video_folder, settings))
        summary.append(
            {
                "account_no": acc.account_no,
                "account_name": acc.account_name,
                "ok": ok,
                "error": err,
                "queue_remaining": pending,
            }
        )
    df_sum = pd.DataFrame(summary)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df_pub.to_excel(writer, sheet_name="Публикации", index=False)
        df_sum.to_excel(writer, sheet_name="Сводка", index=False)
        meta = pd.DataFrame([{"run_id": run_id}])
        meta.to_excel(writer, sheet_name="Meta", index=False)

    wb = load_workbook(path)
    for sheet in ("Публикации", "Сводка"):
        if sheet in wb.sheetnames:
            ws = wb[sheet]
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="DDDDDD")
    wb.save(path)
    return path


def export_last_run(settings: AppSettings, store: StateStore) -> Path | None:
    run_id = store.last_run_id()
    if not run_id:
        return None
    records = store.publications_for_run(run_id)
    out = settings.reports_dir / f"report_{run_id}.xlsx"
    return write_run_report(out, run_id, records, [], settings)
