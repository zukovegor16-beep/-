from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from threading import Lock
from typing import Any

import os
import tempfile

from openpyxl import Workbook, load_workbook

from tiktok_publisher.captions import split_caption_variants
from tiktok_publisher.config import AppSettings

HEADER_ALIASES = {
    "last_run_status": {
        "last_run_status",
        "status",
        "статус",
        "run_status",
    },
    "account_no": {"account_no", "номер", "account", "id"},
    "account_name": {"account_name", "name", "имя", "название"},
    "video_folder": {"video_folder", "folder", "папка", "video_folder_path"},
    "dolphin_profile_id": {"dolphin_profile_id", "profile_id", "dolphin", "profile"},
    "videos_per_run": {"videos_per_run", "videos", "count", "n", "количество"},
    "caption_variants": {"caption_variants", "captions", "подписи", "caption"},
    "active": {"active", "активен", "enabled"},
    "geo": {"geo", "гео", "country", "region"},
}

ONBOARD_HEADER_ALIASES = {
    "geo": {"geo", "гео", "country", "region"},
    "tiktok_login": {"tiktok_login", "login", "логин", "tiktok"},
    "tiktok_password_current": {
        "tiktok_password_current",
        "password_current",
        "password",
        "пароль",
        "текущий_пароль",
    },
    "tiktok_password_new": {
        "tiktok_password_new",
        "password_new",
        "новый_пароль",
    },
    "email": {"email", "почта", "mail"},
    "email_password": {"email_password", "пароль_почты"},
    "avatar_path": {"avatar_path", "avatar", "аватар"},
    "proxy": {"proxy", "прокси"},
    "dolphin_profile_id": {"dolphin_profile_id", "profile_id", "dolphin", "profile"},
    "onboard_status": {"onboard_status", "status", "статус", "onboard"},
    "onboard_error": {"onboard_error", "error", "ошибка"},
}


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower().replace(" ", "_")


def _map_headers(row: tuple[Any, ...], aliases: dict[str, set[str]] | None = None) -> dict[str, int]:
    alias_map = aliases or HEADER_ALIASES
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(row):
        key = _normalize_header(cell)
        for canonical, names in alias_map.items():
            if key == canonical or key in names:
                mapping[canonical] = idx
                break
    return mapping


def _cell(row: tuple[Any, ...], col_map: dict[str, int], name: str) -> Any:
    idx = col_map.get(name)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _parse_active(value: Any) -> bool:
    if value is None:
        return False
    s = str(value).strip().lower()
    return s in {"1", "true", "yes", "y", "да", "on", "active"}


@dataclass
class AccountRow:
    row_index: int  # 1-based Excel row
    account_no: str
    account_name: str
    video_folder: Path
    dolphin_profile_id: str
    videos_per_run: int
    caption_variants_raw: str
    active: bool
    geo: str = ""

    def caption_variants(self, delimiter: str) -> list[str]:
        return split_caption_variants(self.caption_variants_raw, delimiter)


@dataclass
class OnboardRow:
    row_index: int
    geo: str
    tiktok_login: str
    tiktok_password_current: str
    tiktok_password_new: str
    email: str
    email_password: str
    avatar_path: Path
    proxy: str
    dolphin_profile_id: str
    onboard_status: str
    onboard_error: str


def _onboard_done(status: str) -> bool:
    s = (status or "").strip().lower()
    if not s:
        return False
    return s in {x.lower() for x in DONE_ONBOARD_LABELS}


DONE_ONBOARD_LABELS = {"готов", "onboard_ready"}


def _atomic_save_workbook(wb, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(dest.parent))
    os.close(fd)
    tmp_path = Path(tmp)
    try:
        wb.save(tmp_path)
        tmp_path.replace(dest)
    finally:
        if tmp_path.is_file():
            tmp_path.unlink(missing_ok=True)


class ExcelRegistry:
    _write_lock = Lock()

    def __init__(self, settings: AppSettings):
        self.settings = settings
        self.path = settings.excel_path

    def read_accounts(self) -> list[AccountRow]:
        if not self.path.is_file():
            raise FileNotFoundError(f"Excel не найден: {self.path}")
        wb = load_workbook(self.path, read_only=True, data_only=True)
        sheet_name = self.settings.excel.sheet_accounts
        if sheet_name not in wb.sheetnames:
            wb.close()
            raise ValueError(f"Лист «{sheet_name}» не найден в {self.path}")
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        col_map = _map_headers(rows[0])
        required = [
            "account_no",
            "video_folder",
            "dolphin_profile_id",
            "videos_per_run",
            "caption_variants",
            "active",
        ]
        missing = [r for r in required if r not in col_map]
        if missing:
            raise ValueError(f"В Excel нет колонок: {', '.join(missing)}")

        accounts: list[AccountRow] = []
        for i, row in enumerate(rows[1:], start=2):
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
            active = _parse_active(_cell(row, col_map, "active"))
            if not active:
                continue
            try:
                n = int(_cell(row, col_map, "videos_per_run"))
            except (TypeError, ValueError):
                n = 0
            folder_raw = _cell(row, col_map, "video_folder")
            geo_raw = _cell(row, col_map, "geo")
            accounts.append(
                AccountRow(
                    row_index=i,
                    account_no=str(_cell(row, col_map, "account_no")).strip(),
                    account_name=str(_cell(row, col_map, "account_name") or "").strip(),
                    video_folder=Path(str(folder_raw).strip()),
                    dolphin_profile_id=str(_cell(row, col_map, "dolphin_profile_id")).strip(),
                    videos_per_run=max(0, n),
                    caption_variants_raw=str(_cell(row, col_map, "caption_variants") or ""),
                    active=True,
                    geo=str(geo_raw or "").strip(),
                )
            )
        return accounts

    def read_onboarding_rows(self) -> list[OnboardRow]:
        if not self.path.is_file():
            raise FileNotFoundError(f"Excel не найден: {self.path}")
        wb = load_workbook(self.path, read_only=True, data_only=True)
        sheet_name = self.settings.excel.sheet_onboarding
        if sheet_name not in wb.sheetnames:
            wb.close()
            return []
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        col_map = _map_headers(rows[0], ONBOARD_HEADER_ALIASES)
        required = ["geo", "tiktok_login", "tiktok_password_current"]
        missing = [r for r in required if r not in col_map]
        if missing:
            raise ValueError(f"В листе «{sheet_name}» нет колонок: {', '.join(missing)}")

        out: list[OnboardRow] = []
        for i, row in enumerate(rows[1:], start=2):
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
            av = _cell(row, col_map, "avatar_path")
            out.append(
                OnboardRow(
                    row_index=i,
                    geo=str(_cell(row, col_map, "geo") or "").strip(),
                    tiktok_login=str(_cell(row, col_map, "tiktok_login") or "").strip(),
                    tiktok_password_current=str(
                        _cell(row, col_map, "tiktok_password_current") or ""
                    ).strip(),
                    tiktok_password_new=str(
                        _cell(row, col_map, "tiktok_password_new") or ""
                    ).strip(),
                    email=str(_cell(row, col_map, "email") or "").strip(),
                    email_password=str(_cell(row, col_map, "email_password") or "").strip(),
                    avatar_path=Path(str(av).strip()) if av else Path(""),
                    proxy=str(_cell(row, col_map, "proxy") or "").strip(),
                    dolphin_profile_id=str(
                        _cell(row, col_map, "dolphin_profile_id") or ""
                    ).strip(),
                    onboard_status=str(_cell(row, col_map, "onboard_status") or "").strip(),
                    onboard_error=str(_cell(row, col_map, "onboard_error") or "").strip(),
                )
            )
        return out

    def read_onboarding_pending(self) -> list[OnboardRow]:
        return [r for r in self.read_onboarding_rows() if not _onboard_done(r.onboard_status)]

    def update_onboard_row(
        self,
        row_index: int,
        *,
        dolphin_profile_id: str | None = None,
        onboard_status: str | None = None,
        onboard_error: str | None = None,
    ) -> bool:
        if not self.path.is_file():
            return False
        with self._write_lock:
            try:
                wb = load_workbook(self.path)
            except PermissionError:
                return False
            sheet_name = self.settings.excel.sheet_onboarding
            if sheet_name not in wb.sheetnames:
                wb.close()
                return False
            ws = wb[sheet_name]
            headers = [c.value for c in ws[1]]
            col_map = _map_headers(tuple(headers), ONBOARD_HEADER_ALIASES)
            updates: dict[str, Any] = {}
            if dolphin_profile_id is not None:
                updates["dolphin_profile_id"] = dolphin_profile_id
            if onboard_status is not None:
                updates["onboard_status"] = onboard_status
            if onboard_error is not None:
                updates["onboard_error"] = onboard_error
            for field, value in updates.items():
                idx = col_map.get(field)
                if idx is not None:
                    ws.cell(row=row_index, column=idx + 1, value=value)
            try:
                _atomic_save_workbook(wb, self.path)
            except (PermissionError, OSError):
                wb.close()
                return False
            wb.close()
            return True

    def promote_onboard_to_accounts(self, row: OnboardRow, profile_id: str) -> int:
        """Append inactive account row; returns new row index."""
        if not self.path.is_file():
            raise FileNotFoundError(self.path)
        with self._write_lock:
            wb = load_workbook(self.path)
            sheet_name = self.settings.excel.sheet_accounts
            if sheet_name not in wb.sheetnames:
                ws = wb.create_sheet(sheet_name)
                ws.append(
                    [
                        "account_no",
                        "account_name",
                        "video_folder",
                        "dolphin_profile_id",
                        "videos_per_run",
                        "caption_variants",
                        "active",
                        "geo",
                        "status",
                        "last_run_error",
                        "last_run_at",
                    ]
                )
            ws = wb[sheet_name]
            account_no = row.tiktok_login or f"onboard-{row.row_index}"
            ws.append(
                [
                    account_no,
                    row.tiktok_login,
                    "",
                    profile_id,
                    1,
                    "",
                    "нет",
                    row.geo,
                    "",
                    "",
                    "",
                ]
            )
            new_row = ws.max_row
            _atomic_save_workbook(wb, self.path)
            wb.close()
            return new_row

    def update_account_run_status(
        self,
        row_index: int,
        status: str,
        error: str | None = None,
    ) -> bool:
        """Returns False if file locked or missing."""
        if not self.path.is_file():
            return False
        with self._write_lock:
            try:
                wb = load_workbook(self.path)
            except PermissionError:
                return False
            sheet_name = self.settings.excel.sheet_accounts
            if sheet_name not in wb.sheetnames:
                wb.close()
                return False
            ws = wb[sheet_name]
            headers = [c.value for c in ws[1]]
            col_map = _map_headers(tuple(headers))
            now = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
            updates: dict[str, Any] = {
                "last_run_status": status,
                "last_run_error": error or "",
                "last_run_at": now,
            }
            for field, value in updates.items():
                idx = col_map.get(field)
                if idx is not None:
                    ws.cell(row=row_index, column=idx + 1, value=value)
            try:
                _atomic_save_workbook(wb, self.path)
            except PermissionError:
                wb.close()
                return False
            except OSError:
                wb.close()
                return False
            wb.close()
            return True

    def append_log_row(
        self,
        account_no: str,
        filename: str,
        caption_used: str,
        status: str,
        error: str,
        started_at: str,
        finished_at: str,
    ) -> bool:
        if not self.path.is_file():
            return False
        with self._write_lock:
            try:
                wb = load_workbook(self.path)
            except PermissionError:
                return False
            log_name = "Log"
            if log_name not in wb.sheetnames:
                ws = wb.create_sheet(log_name)
                ws.append(
                    [
                        "started_at",
                        "finished_at",
                        "account_no",
                        "file",
                        "caption_used",
                        "status",
                        "error",
                    ]
                )
            ws = wb[log_name]
            ws.append(
                [
                    started_at,
                    finished_at,
                    account_no,
                    filename,
                    caption_used,
                    status,
                    error,
                ]
            )
            try:
                _atomic_save_workbook(wb, self.path)
            except PermissionError:
                wb.close()
                return False
            except OSError:
                wb.close()
                return False
            wb.close()
            return True


def create_workbook_from_template(template: Path, dest: Path) -> None:
    if dest.is_file():
        return
    if not template.is_file():
        wb = Workbook()
        ws = wb.active
        ws.title = "Аккаунты"
        ws.append(
            [
                "account_no",
                "account_name",
                "video_folder",
                "dolphin_profile_id",
                "videos_per_run",
                "caption_variants",
                "active",
                "geo",
                "status",
                "last_run_error",
                "last_run_at",
            ]
        )
        ob = wb.create_sheet("Онбординг")
        ob.append(
            [
                "geo",
                "tiktok_login",
                "tiktok_password_current",
                "tiktok_password_new",
                "email",
                "email_password",
                "avatar_path",
                "proxy",
                "dolphin_profile_id",
                "onboard_status",
                "onboard_error",
            ]
        )
        dest.parent.mkdir(parents=True, exist_ok=True)
        wb.save(dest)
        return
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(template.read_bytes())
